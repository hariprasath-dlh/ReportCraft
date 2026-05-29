"""
Full integration test: generate a complete PPT and validate output.
"""
import sys
import os
from pathlib import Path

# Fix Windows console encoding
if sys.platform == "win32":
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    sys.stderr.reconfigure(encoding='utf-8', errors='replace')

sys.path.insert(0, str(Path(__file__).parent))

os.environ.setdefault("OUTPUT_DIR", str(Path(__file__).parent / "outputs"))
os.environ.setdefault("UPLOAD_DIR", str(Path(__file__).parent / "uploads"))

from models.schemas import (
    PPTRequest, TitleSlide, Student, LiteratureEntry, StyleInstructions
)
from generators.ppt_generator import generate_ppt


def make_review0_request():
    return PPTRequest(
        project_type="mini_project_2",
        review_type="review_0",
        title_slide=TitleSlide(
            course_code="22UAD607",
            course_name="MINI PROJECT-II",
            project_name="Multilingual Campus Assistant",
            students=[
                Student(name="ABHINANDHANA K V", reg_no="23AD002"),
                Student(name="D NAVANEESH", reg_no="23AD023"),
            ],
            guide_name="Dr. I.DEVI",
            guide_designation="ASSISTANT PROFESSOR",
            department="Artificial Intelligence and Data Science",
        ),
        abstract=(
            "This project presents a Multilingual Campus Assistant that helps "
            "students and visitors navigate the college campus with real-time "
            "translation support. The system uses natural language processing "
            "to provide guidance in multiple languages, making the campus "
            "accessible to a diverse student body."
        ),
        problem_statement=(
            "Many campus visitors and newly enrolled students from diverse "
            "linguistic backgrounds find it difficult to navigate the college "
            "campus, understand signage, and access information in their "
            "preferred language."
        ),
        existing_system=(
            "The current system relies on physical signboards in English/Tamil "
            "and manual guidance from security personnel. This approach is "
            "limited to certain hours and cannot handle dynamic queries."
        ),
        proposed_system=(
            "We propose an AI-powered multilingual assistant that uses NLP "
            "and speech recognition to provide campus navigation, event info, "
            "and administrative guidance in 10+ languages via a mobile app."
        ),
        instructions=StyleInstructions(),
    )


def make_review1_request():
    return PPTRequest(
        project_type="mini_project_2",
        review_type="review_1",
        title_slide=TitleSlide(
            course_code="22UAD607",
            course_name="MINI PROJECT-II",
            project_name="Multilingual Campus Assistant",
            students=[
                Student(name="ABHINANDHANA K V", reg_no="23AD002"),
                Student(name="D NAVANEESH", reg_no="23AD023"),
                Student(name="PRIYA SHARMA", reg_no="23AD045"),
            ],
            guide_name="Dr. I.DEVI",
            guide_designation="ASSISTANT PROFESSOR",
            department="Artificial Intelligence and Data Science",
        ),
        abstract="This project presents a Multilingual Campus Assistant.",
        problem_statement="Many campus visitors face language barriers.",
        existing_system="Current system relies on physical signboards.",
        proposed_system="AI-powered multilingual assistant using NLP.",
        literature_survey=[
            LiteratureEntry(
                title="Multilingual NLP in Education Settings",
                author="Kumar, A. and Zhang, L.",
                year="2023",
                summary="Explores NLP techniques for multilingual support in university environments."
            ),
            LiteratureEntry(
                title="Campus Navigation Systems: A Survey",
                author="Johnson, M. et al.",
                year="2022",
                summary="Reviews indoor navigation systems using beacons, Wi-Fi, and AR."
            ),
            LiteratureEntry(
                title="Speech Recognition for Indian Languages",
                author="Patel, R. and Devi, S.",
                year="2023",
                summary="Surveys ASR models for Hindi, Tamil, and Telugu."
            ),
            LiteratureEntry(
                title="Chatbot Design Patterns for Higher Education",
                author="Smith, J. and Brown, K.",
                year="2021",
                summary="Identifies effective chatbot architectures for campus use."
            ),
            LiteratureEntry(
                title="Mobile-First UX for Accessibility",
                author="Chen, W.",
                year="2024",
                summary="Discusses mobile accessibility patterns for diverse user populations."
            ),
        ],
        instructions=StyleInstructions(),
    )


def run_test(name, request):
    print(f"\n{'='*60}")
    print(f"TEST: {name}")
    print(f"{'='*60}")
    try:
        file_id, path, total_slides = generate_ppt(request)
        print(f"  [OK] Generated successfully")
        print(f"  [OK] File: {path}")
        print(f"  [OK] File exists: {path.exists()}")
        print(f"  [OK] File size: {path.stat().st_size:,} bytes")
        print(f"  [OK] Total slides: {total_slides}")

        # Validate the PPTX can be re-opened
        from pptx import Presentation
        prs = Presentation(str(path))
        actual_slides = len(prs.slides)
        print(f"  [OK] Re-opened successfully: {actual_slides} slides")
        assert actual_slides == total_slides, \
            f"Expected {total_slides} slides but got {actual_slides}"
        print(f"  [OK] Slide count matches")

        # Check slide dimensions
        from pptx.util import Inches
        assert prs.slide_width == Inches(10), "Slide width should be 10 inches"
        assert prs.slide_height == Inches(7.5), "Slide height should be 7.5 inches"
        print(f"  [OK] Slide dimensions: 10 x 7.5 inches (4:3)")

        # Validate each slide has shapes
        for i, slide in enumerate(prs.slides):
            shape_count = len(slide.shapes)
            print(f"    Slide {i+1}: {shape_count} shapes")
            assert shape_count > 0, f"Slide {i+1} has no shapes!"

        print(f"\n  PASSED: {name}")
        return True

    except Exception as e:
        print(f"\n  FAILED: {type(e).__name__}: {e}")
        import traceback
        traceback.print_exc()
        return False


if __name__ == "__main__":
    results = {}
    results["Review 0"] = run_test("Review 0 (basic)", make_review0_request())
    results["Review 1"] = run_test("Review 1 (with literature)", make_review1_request())

    print(f"\n{'='*60}")
    print("SUMMARY")
    print(f"{'='*60}")
    for name, passed in results.items():
        status = "PASS" if passed else "FAIL"
        print(f"  [{status}] {name}")

    all_pass = all(results.values())
    print(f"\nOverall: {'ALL PASSED' if all_pass else 'SOME FAILED'}")
    sys.exit(0 if all_pass else 1)
