"""
Standalone PPT smoke test -- verifies python-pptx and all helpers work.
Run: python test_pptx_basic.py  (from backend/ directory)
No backend server needed. No internet needed.
"""
import sys
from pathlib import Path

# Fix Windows console encoding
if sys.platform == "win32":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

passed = 0
failed = 0


def check(label, fn):
    global passed, failed
    try:
        fn()
        print(f"  [OK] {label}")
        passed += 1
    except Exception as e:
        print(f"  [FAIL] {label}: {type(e).__name__}: {e}")
        failed += 1


# ── Test 1: Core imports ─────────────────────────────────────────────────
print("\n== 1. Core Imports ==")

def test_pptx_import():
    from pptx import Presentation  # noqa: F401

def test_pptx_util():
    from pptx.util import Inches, Pt  # noqa: F401

def test_pptx_color():
    from pptx.dml.color import RGBColor  # noqa: F401

def test_pptx_enum_text():
    from pptx.enum.text import PP_ALIGN  # noqa: F401

def test_pptx_enum_shapes():
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR  # noqa: F401

def test_pptx_oxml():
    from pptx.oxml.ns import qn  # noqa: F401

def test_lxml():
    from lxml import etree  # noqa: F401

def test_pillow():
    from PIL import Image  # noqa: F401

check("python-pptx core", test_pptx_import)
check("pptx.util (Inches, Pt)", test_pptx_util)
check("pptx.dml.color (RGBColor)", test_pptx_color)
check("pptx.enum.text (PP_ALIGN)", test_pptx_enum_text)
check("pptx.enum.shapes (MSO_SHAPE)", test_pptx_enum_shapes)
check("pptx.oxml.ns (qn)", test_pptx_oxml)
check("lxml (etree)", test_lxml)
check("Pillow (PIL)", test_pillow)


# ── Test 2: Create a basic presentation ──────────────────────────────────
print("\n== 2. Basic Presentation ==")

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

prs = None

def test_create_prs():
    global prs
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    assert prs.slide_width == Inches(10)
    assert prs.slide_height == Inches(7.5)

check("Create presentation (10x7.5)", test_create_prs)


# ── Test 3: Add a blank slide ────────────────────────────────────────────
print("\n== 3. Blank Slide ==")

slide = None

def test_add_slide():
    global slide
    try:
        layout = prs.slide_layouts[6]
    except IndexError:
        layout = prs.slide_layouts[-1]
    slide = prs.slides.add_slide(layout)
    assert slide is not None

check("Add blank slide", test_add_slide)


# ── Test 4: Add navy bar shape ───────────────────────────────────────────
print("\n== 4. Navy Bar Shape ==")

def test_navy_bar():
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0), Inches(0),
        Inches(1.18), Inches(7.5)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor(0x1F, 0x38, 0x64)
    bar.line.fill.background()
    assert bar is not None

check("Add navy rectangle shape", test_navy_bar)


# ── Test 5: Add text ─────────────────────────────────────────────────────
print("\n== 5. Text Box ==")

def test_text_box():
    txb = slide.shapes.add_textbox(
        Inches(1.5), Inches(2), Inches(7), Inches(3)
    )
    tf = txb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = "Hello ReportCraft"
    r.font.name = "Times New Roman"
    r.font.size = Pt(24)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 0, 0)
    assert r.text == "Hello ReportCraft"

check("Add text box with run", test_text_box)


# ── Test 6: Save and reopen ──────────────────────────────────────────────
print("\n== 6. Save & Reopen ==")

out_path = Path("test_smoke.pptx")

def test_save():
    prs.save(str(out_path))
    assert out_path.exists()
    assert out_path.stat().st_size > 0

def test_reopen():
    prs2 = Presentation(str(out_path))
    assert len(prs2.slides) == 1
    assert prs2.slide_width == Inches(10)

def test_cleanup():
    out_path.unlink()
    assert not out_path.exists()

check("Save to test_smoke.pptx", test_save)
check("Reopen and validate", test_reopen)
check("Cleanup test file", test_cleanup)


# ── Test 7: College logo exists ──────────────────────────────────────────
print("\n== 7. College Logo ==")

LOGO_PATHS = [
    Path("templates/ppt/college_logo.png"),
    Path("templates/ppt/college_logo.jpg"),
    Path("templates/ppt/college_logo.jpeg"),
]

def test_logo_exists():
    found = any(p.exists() for p in LOGO_PATHS)
    if not found:
        existing = [str(p) for p in Path("templates/ppt").iterdir()] if Path("templates/ppt").exists() else []
        raise FileNotFoundError(
            f"No college logo found. Checked: {[str(p) for p in LOGO_PATHS]}. "
            f"Files in templates/ppt/: {existing}"
        )

check("College logo file present", test_logo_exists)


# ── Test 8: Internal helpers import ──────────────────────────────────────
print("\n== 8. ReportCraft Helpers ==")

# Add backend to path for imports
sys.path.insert(0, str(Path(__file__).parent))

def test_helpers_import():
    from generators.slide_helpers import (
        add_navy_bar, add_footer, add_heading,
        add_content_box, set_line_spacing,
        set_slide_background_white, parse_numbered_or_bullet,
    )

def test_builders_import():
    from generators.slide_builder import (
        build_title_slide, build_abstract_slide,
        build_problem_statement_slide, build_literature_slide,
        build_existing_system_slide, build_proposed_system_slide,
    )

def test_schemas_import():
    from models.schemas import PPTRequest, TitleSlide, Student

check("slide_helpers imports", test_helpers_import)
check("slide_builder imports", test_builders_import)
check("schemas imports", test_schemas_import)


# ── Summary ──────────────────────────────────────────────────────────────
print(f"\n{'='*50}")
print(f"RESULTS: {passed} passed, {failed} failed")
print(f"{'='*50}")

if failed:
    print("\nFix the failed items before running the full generator.")
    sys.exit(1)
else:
    print("\nAll checks passed -- environment is ready for PPT generation.")
    sys.exit(0)
