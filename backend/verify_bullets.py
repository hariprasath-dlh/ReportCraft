import sys
import os
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches

# Add backend directory to path
sys.path.insert(0, str(Path(__file__).parent))

from models.schemas import PPTRequest, TitleSlide, Student, LiteratureEntry, StyleInstructions, PPTSection
from generators.ppt_generator import generate_ppt

def run_verification():
    print("=== STARTING AUTOMATED CHECKLIST VERIFICATION ===")

    # Setup college logo dummy if not exists
    logo_path = Path("templates/ppt/college_logo.png")
    if not logo_path.exists():
        logo_path.parent.mkdir(parents=True, exist_ok=True)
        from PIL import Image
        img = Image.new('RGB', (100, 100), color = 'blue')
        img.save(logo_path)

    # 1. Build test request representing Test A through Test G
    request = PPTRequest(
        project_type="mini_project_2",
        review_type="review_1",
        title_slide=TitleSlide(
            course_code="22UAD607",
            course_name="MINI PROJECT-II",
            project_name="Bullet Format Verification",
            students=[Student(name="VERIFICATION BOT", reg_no="CHECKLIST")],
            guide_name="TEST GUIDE",
            guide_designation="PROFESSOR",
            department="CSE",
        ),
        literature_survey=[
            LiteratureEntry(title=f"Lit {i}", author="Author", year="2023", summary="Summary")
            for i in range(5)
        ],
        instructions=StyleInstructions(),
        # Configure sections to test cases
        sections=[
            # Test A: format_as_bullets = False, paragraph content
            PPTSection(
                key="abstract",
                title="Test A: Paragraph Style",
                content="This is a standard paragraph content. It should not render as bullet points since the format flag is set to false.",
                is_enabled=True,
                order=0,
                format_as_bullets=False
            ),
            # Test B: format_as_bullets = True, paragraph content (single long line to split)
            # Must exceed 150 characters to trigger splitting.
            PPTSection(
                key="problem",
                title="Test B: Paragraph Split to Bullets",
                content="This is the first sentence in our paragraph to exceed length. This is the second sentence of the slide that we want to split! And this is the third sentence of the slide to test sentence boundary parsing.",
                is_enabled=True,
                order=1,
                format_as_bullets=True
            ),
            # Test G: Disabled section with format_as_bullets = True (should not render)
            PPTSection(
                key="existing",
                title="Test G: Disabled Slide",
                content="This section is disabled. It should not appear in the final PPT slides at all.",
                is_enabled=False,
                order=2,
                format_as_bullets=True
            ),
            # Test E: Checkbox ON but content is empty (should skip content box)
            PPTSection(
                key="proposed",
                title="Test E: Empty Content Bullets",
                content="",
                is_enabled=True,
                order=3,
                format_as_bullets=True
            ),
        ],
        custom_sections=[
            # Test C: Custom section, checkbox ON, user typed "- item1\n- item2"
            PPTSection(
                key="custom_c",
                title="Test C: Clean Dash Bullets",
                content="- Item number one\n- Item number two\n* Item number three",
                is_enabled=True,
                order=4,
                format_as_bullets=True
            ),
            # Test D: Custom section, checkbox ON, user typed "1. a\n2. b\n3. c"
            PPTSection(
                key="custom_d",
                title="Test D: Clean Number Bullets",
                content="1. First item\n2) Second item\n3. Third item",
                is_enabled=True,
                order=5,
                format_as_bullets=True
            ),
            # Test F: One checkbox ON, one OFF (this file has a mix of ON and OFF; we will verify custom_off is OFF)
            PPTSection(
                key="custom_off",
                title="Test F: Custom Off Paragraph",
                content="This is a custom section with bullets toggle OFF. It should render as a paragraph.",
                is_enabled=True,
                order=6,
                format_as_bullets=False
            ),
            # Case 1: Checkbox ticked but content is a single short line
            PPTSection(
                key="custom_short",
                title="Case 1: Single Short Line",
                content="This is a single short line.",
                is_enabled=True,
                order=7,
                format_as_bullets=True
            )
        ]
    )

    # 2. Run generator
    print("Generating verification PPT...")
    file_id, output_path, total_slides = generate_ppt(
        request,
        image_paths=[],
        architecture_path=None,
        section_image_map={}
    )
    print(f"Generated PPTX file: {output_path}")

    # 3. Read and verify slides
    prs = Presentation(str(output_path))
    
    # Slides mapping:
    # 1. Title
    # 2. Test A: Paragraph Style
    # 3. Test B: Paragraph Split to Bullets
    # 4-8. Literature Survey (5 slides)
    # 9. Test E: Empty Content Bullets
    # 10. Test C: Clean Dash Bullets
    # 11. Test D: Clean Number Bullets
    # 12. Test F: Custom Off Paragraph
    # 13. Case 1: Single Short Line
    # Total expected slides: 13 (disabled slide is excluded)
    print(f"Total slides: {len(prs.slides)}. Expected: 13.")
    assert len(prs.slides) == 13, f"Expected 13 slides, got {len(prs.slides)}"
    print("[OK] Test G passed: Disabled section is successfully excluded (slide count matches expectations).")

    # Let's map slides by their headings to verify their content box structure
    slide_by_title = {}
    for slide in prs.slides:
        title_text = ""
        # The title is usually a textbox in the slide. Let's find it.
        # Slide headings are typically added via add_heading which puts a textbox with a single paragraph
        for shape in slide.shapes:
            if shape.has_text_frame:
                tf = shape.text_frame
                text = tf.text.strip()
                if "Test" in text or "Case" in text:
                    title_text = text
                    break
        if title_text:
            slide_by_title[title_text] = slide

    # Spatial content box filter helper
    def get_content_box(slide):
        # Excludes left bar (left is negative or < 1.0 inches), header (top < 0.9 inches), footer (top > 7.0 inches)
        boxes = []
        for s in slide.shapes:
            if s.has_text_frame:
                # Inches conversion for comparison
                left_in = s.left / 914400.0
                top_in = s.top / 914400.0
                if left_in >= 1.0 and 0.9 <= top_in <= 7.0:
                    # Double-check it's not the heading title itself
                    # (heading is sometimes at top_in = 0.3 but let's be double sure it's not the slide title text)
                    if s.text_frame.text.strip() not in slide_by_title.keys():
                        boxes.append(s)
        return boxes

    # Let's verify each test case
    
    # Test A: Paragraph Style (format_as_bullets = False)
    # Should render as a paragraph (1 paragraph in content box, not starting with •)
    slide_a = slide_by_title.get("Test A: Paragraph Style")
    assert slide_a is not None, "Could not find Slide A"
    content_boxes = get_content_box(slide_a)
    assert len(content_boxes) == 1, f"Expected 1 content box, got {len(content_boxes)}"
    tf = content_boxes[0].text_frame
    paragraphs = [p.text for p in tf.paragraphs if p.text.strip()]
    assert len(paragraphs) == 1, f"Expected 1 paragraph, got {len(paragraphs)}"
    assert not paragraphs[0].startswith("•"), f"Paragraph should not start with bullet marker: {paragraphs[0]}"
    print("[OK] Test A passed: Slide rendered as standard paragraph.")

    # Test B: Paragraph Split to Bullets (format_as_bullets = True)
    # Should split single long line by sentence boundaries and render as 3 bullets
    slide_b = slide_by_title.get("Test B: Paragraph Split to Bullets")
    assert slide_b is not None, "Could not find Slide B"
    content_boxes = get_content_box(slide_b)
    assert len(content_boxes) == 1, f"Expected 1 content box, got {len(content_boxes)}"
    tf = content_boxes[0].text_frame
    paragraphs = [p.text.strip() for p in tf.paragraphs if p.text.strip()]
    assert len(paragraphs) == 3, f"Expected 3 bullet paragraphs, got {len(paragraphs)}"
    for p in paragraphs:
        assert p.startswith("•"), f"Bullet paragraph should start with standard bullet symbol: {p}"
    assert "This is the first sentence in our paragraph to exceed length." in paragraphs[0]
    assert "This is the second sentence of the slide that we want to split!" in paragraphs[1]
    assert "And this is the third sentence of the slide to test sentence boundary parsing." in paragraphs[2]
    print("[OK] Test B passed: Paragraph successfully split into 3 clean bullet points.")

    # Test E: Checkbox ON but content is empty
    # Should skip content box entirely
    slide_e = slide_by_title.get("Test E: Empty Content Bullets")
    assert slide_e is not None, "Could not find Slide E"
    content_boxes = get_content_box(slide_e)
    assert len(content_boxes) == 0, f"Expected 0 content boxes for empty slide, got {len(content_boxes)}"
    print("[OK] Test E passed: Empty bullet slide has no content text box.")

    # Test C: Custom section, checkbox ON, user typed "- item1\n- item2"
    # Should clean dash/star markers and show bullets without doubles
    slide_c = slide_by_title.get("Test C: Clean Dash Bullets")
    assert slide_c is not None, "Could not find Slide C"
    content_boxes = get_content_box(slide_c)
    assert len(content_boxes) == 1, f"Expected 1 content box, got {len(content_boxes)}"
    tf = content_boxes[0].text_frame
    paragraphs = [p.text.strip() for p in tf.paragraphs if p.text.strip()]
    assert len(paragraphs) == 3, f"Expected 3 bullets, got {len(paragraphs)}"
    assert paragraphs[0] == "•  Item number one", f"Expected clean bullet text, got: {paragraphs[0]}"
    assert paragraphs[1] == "•  Item number two"
    assert paragraphs[2] == "•  Item number three"
    print("[OK] Test C passed: Leading dashes and stars stripped to prevent double bullets.")

    # Test D: Custom section, checkbox ON, user typed "1. a\n2. b\n3. c"
    # Should clean numbers and show bullets
    slide_d = slide_by_title.get("Test D: Clean Number Bullets")
    assert slide_d is not None, "Could not find Slide D"
    content_boxes = get_content_box(slide_d)
    assert len(content_boxes) == 1, f"Expected 1 content box, got {len(content_boxes)}"
    tf = content_boxes[0].text_frame
    paragraphs = [p.text.strip() for p in tf.paragraphs if p.text.strip()]
    assert len(paragraphs) == 3, f"Expected 3 bullets, got {len(paragraphs)}"
    assert paragraphs[0] == "•  First item"
    assert paragraphs[1] == "•  Second item"
    assert paragraphs[2] == "•  Third item"
    print("[OK] Test D passed: Leading list numbers cleaned and rendered as standard bullets.")

    # Test F: One checkbox ON, another OFF in same PPT
    # We verified multiple ON above, now verify custom_off is OFF (paragraph)
    slide_f = slide_by_title.get("Test F: Custom Off Paragraph")
    assert slide_f is not None, "Could not find Slide F"
    content_boxes = get_content_box(slide_f)
    assert len(content_boxes) == 1, f"Expected 1 content box, got {len(content_boxes)}"
    tf = content_boxes[0].text_frame
    paragraphs = [p.text.strip() for p in tf.paragraphs if p.text.strip()]
    assert len(paragraphs) == 1, f"Expected 1 paragraph, got {len(paragraphs)}"
    assert not paragraphs[0].startswith("•"), "Should render as paragraph, not bullets"
    print("[OK] Test F passed: Mixed ON and OFF sections in same PPT behave correctly per section.")

    # Case 1: Checkbox ticked but content is a single short line
    # Should render as a single bullet point
    slide_short = slide_by_title.get("Case 1: Single Short Line")
    assert slide_short is not None, "Could not find Case 1 Slide"
    content_boxes = get_content_box(slide_short)
    assert len(content_boxes) == 1, f"Expected 1 content box, got {len(content_boxes)}"
    tf = content_boxes[0].text_frame
    paragraphs = [p.text.strip() for p in tf.paragraphs if p.text.strip()]
    assert len(paragraphs) == 1, f"Expected 1 bullet, got {len(paragraphs)}"
    assert paragraphs[0] == "•  This is a single short line.", f"Expected bullet for short line, got {paragraphs[0]}"
    print("[OK] Case 1 passed: Single short line successfully rendered as bullet.")

    print("\nALL PPT GENERATION LAYOUT ASSERTIONS PASSED SUCCESSFULLY!")

if __name__ == "__main__":
    run_verification()
