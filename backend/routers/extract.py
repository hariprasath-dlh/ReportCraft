"""
Extract routes — POST /api/extract/ppt and POST /api/extract/report
Extract text content from existing presentation/report files.
"""

from io import BytesIO
from fastapi import APIRouter, HTTPException, UploadFile, File
from models.responses import ExtractResponse
from utils.logger import logger

router = APIRouter()


@router.post("/ppt", response_model=ExtractResponse)
async def extract_from_ppt(file: UploadFile = File(...)):
    """Extract text content from an existing PowerPoint file."""
    try:
        from pptx import Presentation
        logger.info(f"Extracting from PPT: {file.filename}")
        content = await file.read()
        prs = Presentation(BytesIO(content))
        sections = {}
        for i, slide in enumerate(prs.slides):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    text = shape.text_frame.text.strip()
                    if text and text != "Dr. NGP INSTITUTE OF TECHNOLOGY":
                        texts.append(text)
            if texts:
                sections[f"slide_{i + 1}"] = " ".join(texts)
        return ExtractResponse(
            sections=sections,
            source_filename=file.filename or "",
        )
    except Exception as e:
        logger.error(f"PPT extraction failed: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/report", response_model=ExtractResponse)
async def extract_from_report(file: UploadFile = File(...)):
    """Extract text content from an existing Word report."""
    try:
        from docx import Document
        logger.info(f"Extracting from report: {file.filename}")
        content = await file.read()
        doc = Document(BytesIO(content))
        full_text = "\n".join(
            p.text for p in doc.paragraphs if p.text.strip()
        )
        return ExtractResponse(
            sections={"full_text": full_text},
            source_filename=file.filename or "",
        )
    except Exception as e:
        logger.error(f"Report extraction failed: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=str(e))
