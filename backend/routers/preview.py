"""
Preview routes — HTML preview for reports and slide list for PPTs.
"""

from pathlib import Path

from fastapi import APIRouter, HTTPException
from pptx import Presentation

from models.schemas import PreviewRequest
from models.responses import PreviewResponse, PPTPreviewResponse
from utils.file_manager import get_output_path
from utils.logger import logger

router = APIRouter()


@router.post("/report", response_model=PreviewResponse)
async def preview_report(request: PreviewRequest):
    """Generate an HTML preview of a DOCX report using mammoth."""
    try:
        file_path = get_output_path(request.file_id, ".docx")
        if not file_path.exists():
            raise HTTPException(
                status_code=404,
                detail=f"Report file not found: {request.file_id}"
            )

        import mammoth
        with open(str(file_path), "rb") as f:
            result = mammoth.convert_to_html(f)

        return PreviewResponse(html=result.value)

    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Report preview failed: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/ppt", response_model=PPTPreviewResponse)
async def preview_ppt(request: PreviewRequest):
    """Generate a slide summary for a PPT file."""
    try:
        file_path = get_output_path(request.file_id, ".pptx")
        if not file_path.exists():
            raise HTTPException(
                status_code=404,
                detail=f"PPT file not found: {request.file_id}"
            )

        prs = Presentation(str(file_path))
        slides = []
        for i, slide in enumerate(prs.slides):
            # Try to find a title
            title = f"Slide {i + 1}"
            if slide.shapes.title:
                title = slide.shapes.title.text or title

            # Collect text from non-title shapes
            texts = []
            for sh in slide.shapes:
                if sh.has_text_frame and sh != getattr(slide.shapes, 'title', None):
                    text = sh.text_frame.text.strip()
                    if text and text != "Dr. NGP INSTITUTE OF TECHNOLOGY":
                        texts.append(text)

            summary = " ".join(texts)[:150]
            slides.append({
                "number": i + 1,
                "title": title,
                "summary": summary,
            })

        return PPTPreviewResponse(
            slides=slides,
            total_slides=len(slides),
        )

    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"PPT preview failed: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=str(e))
