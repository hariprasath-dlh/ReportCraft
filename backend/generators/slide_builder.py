"""
Slide Builder — Per-slide builder functions for the college PPT template.
Each function creates one complete slide with navy bar, footer, and content.
Uses slide_helpers for all reusable layout components.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_CONNECTOR
from pathlib import Path

from generators.slide_helpers import (
    FONT, NAVY, BLACK, WHITE,
    CONTENT_LEFT, CONTENT_WIDTH, SLIDE_W, SLIDE_H,
    add_navy_bar, add_footer, add_heading, add_heading_centre,
    add_content_box, add_content_with_bullets, add_styled_textbox,
    add_image_centred, set_line_spacing, set_space_after,
    set_slide_background_white, parse_numbered_or_bullet,
    get_body_font_size,
)
from utils.logger import logger
from utils.text_formatter import content_to_bullet_points

# Logo path — relative to backend/ working directory
LOGO_PATHS = (
    Path("templates/ppt/college_logo.png"),
    Path("templates/ppt/college_logo.jpg"),
    Path("templates/ppt/college_logo.jpeg"),
)

# Title-slide navy blue per the official template spec
TITLE_NAVY = RGBColor(0x1F, 0x3A, 0x8A)

REVIEW_TYPE_LABELS = {
    "review_0": "ZEROTH REVIEW",
    "review_1": "FIRST REVIEW",
    "review_2": "SECOND REVIEW",
    "review_3": "THIRD REVIEW",
    "review_4": "FOURTH REVIEW",
    "final": "FINAL REVIEW",
}


def _new_blank_slide(prs):
    """Add a new blank slide (layout index 6 = blank)."""
    # Try blank layout; fallback to first available
    try:
        layout = prs.slide_layouts[6]
    except IndexError:
        layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    set_slide_background_white(slide)
    return slide


def _add_centered_textbox(slide, text, top, left, width, height,
                           font_size, bold=False, italic=False,
                           color=RGBColor(0, 0, 0), font_name="Calibri"):
    """Adds a horizontally + vertically centered textbox."""
    txb = slide.shapes.add_textbox(left, top, width, height)
    tf = txb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = Inches(0)
    tf.margin_right = Inches(0)
    tf.margin_top = Inches(0)
    tf.margin_bottom = Inches(0)

    # Vertical anchor — middle
    from pptx.oxml.ns import qn
    txBody = txb._element.find(qn('p:txBody'))
    if txBody is not None:
        bodyPr = txBody.find(qn('a:bodyPr'))
        if bodyPr is not None:
            bodyPr.set('anchor', 'ctr')

    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.name = font_name
    r.font.size = Pt(font_size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    return txb


# ═══════════════════════════════════════════════════════════════════════════
# TITLE SLIDE
# ═══════════════════════════════════════════════════════════════════════════

def build_title_slide(prs, data, slide_num: int):
    """Build the title slide with logo, course info, students, and guide.

    Layout matches the official Dr. NGP Institute of Technology template (4:3 format).
    """
    slide = _new_blank_slide(prs)

    # Clear any existing placeholder shapes on the slide
    for ph in list(slide.placeholders):
        sp = ph._element
        sp.getparent().remove(sp)
    logger.info("Title slide: cleared existing placeholders")

    # ── Left Navy Bar and Footer ──────────────────────────────────────────
    add_navy_bar(slide)
    add_footer(slide, data.title_slide.department, slide_num)

    # ── Logo (top-left of content area) ───────────────────────────────────
    logo_path = next((path for path in LOGO_PATHS if path.exists()), None)
    if logo_path:
        try:
            slide.shapes.add_picture(
                str(logo_path),
                Inches(1.465), Inches(0.41),
                Inches(1.279), Inches(1.195)
            )
            logger.info("Title slide: added logo")
        except Exception as e:
            logger.warning(f"Title slide: could not add logo: {e}")

    # ── Course / Review Details (Top-Right) ───────────────────────────────
    course_txb = slide.shapes.add_textbox(
        Inches(2.8), Inches(0.41), Inches(5.8), Inches(1.2)
    )
    course_tf = course_txb.text_frame
    course_tf.word_wrap = True
    course_tf.margin_left = Inches(0)
    course_tf.margin_right = Inches(0)
    course_tf.margin_top = Inches(0)
    course_tf.margin_bottom = Inches(0)

    p_course = course_tf.paragraphs[0]
    p_course.alignment = PP_ALIGN.CENTER
    r_course = p_course.add_run()
    r_course.text = f"{data.title_slide.course_code}- {data.title_slide.course_name.upper()}"
    r_course.font.name = FONT
    r_course.font.size = Pt(18)
    r_course.font.bold = True
    r_course.font.color.rgb = BLACK
    set_line_spacing(p_course, 115)

    p_rev = course_tf.add_paragraph()
    p_rev.alignment = PP_ALIGN.CENTER
    r_rev = p_rev.add_run()
    r_rev.text = REVIEW_TYPE_LABELS.get(data.review_type, data.review_type).upper()
    r_rev.font.name = FONT
    r_rev.font.size = Pt(18)
    r_rev.font.bold = True
    r_rev.font.color.rgb = BLACK
    set_line_spacing(p_rev, 115)

    # ── Project Title (Center) ────────────────────────────────────────────
    title_txb = slide.shapes.add_textbox(
        Inches(1.38), Inches(2.0), Inches(8.02), Inches(1.5)
    )
    title_tf = title_txb.text_frame
    title_tf.word_wrap = True
    title_tf.margin_left = Inches(0)
    title_tf.margin_right = Inches(0)
    title_tf.margin_top = Inches(0)
    title_tf.margin_bottom = Inches(0)

    # Set vertical anchor to middle
    from pptx.oxml.ns import qn
    txBody = title_txb._element.find(qn('p:txBody'))
    if txBody is not None:
        bodyPr = txBody.find(qn('a:bodyPr'))
        if bodyPr is not None:
            bodyPr.set('anchor', 'ctr')

    p_title = title_tf.paragraphs[0]
    p_title.alignment = PP_ALIGN.CENTER
    r_title = p_title.add_run()
    r_title.text = data.title_slide.project_name
    r_title.font.name = FONT
    r_title.font.size = Pt(24)
    r_title.font.bold = True
    r_title.font.color.rgb = BLACK
    set_line_spacing(p_title, 115)

    # ── Left Column: "Presented by," + Student names and reg numbers ──────
    pres_txb = slide.shapes.add_textbox(
        Inches(1.45), Inches(3.7), Inches(4.6), Inches(0.35)
    )
    pres_tf = pres_txb.text_frame
    pres_tf.word_wrap = True
    pres_tf.margin_left = Inches(0)
    pres_tf.margin_right = Inches(0)
    pres_tf.margin_top = Inches(0)
    pres_tf.margin_bottom = Inches(0)

    p_pres = pres_tf.paragraphs[0]
    p_pres.alignment = PP_ALIGN.LEFT
    r_pres = p_pres.add_run()
    r_pres.text = "Presented by,"
    r_pres.font.name = FONT
    r_pres.font.size = Pt(18)
    r_pres.font.bold = False
    r_pres.font.italic = True
    r_pres.font.color.rgb = BLACK

    # Student names (TextBox A)
    student_name_txb = slide.shapes.add_textbox(
        Inches(1.505), Inches(4.167), Inches(2.9), Inches(2.2)
    )
    student_name_tf = student_name_txb.text_frame
    student_name_tf.word_wrap = True
    student_name_tf.margin_left = Inches(0)
    student_name_tf.margin_right = Inches(0)
    student_name_tf.margin_top = Inches(0)
    student_name_tf.margin_bottom = Inches(0)

    # Student register numbers (TextBox B)
    student_reg_txb = slide.shapes.add_textbox(
        Inches(4.45), Inches(4.167), Inches(1.6), Inches(2.2)
    )
    student_reg_tf = student_reg_txb.text_frame
    student_reg_tf.word_wrap = True
    student_reg_tf.margin_left = Inches(0)
    student_reg_tf.margin_right = Inches(0)
    student_reg_tf.margin_top = Inches(0)
    student_reg_tf.margin_bottom = Inches(0)

    for i, s in enumerate(data.title_slide.students):
        # Name
        p_name = student_name_tf.paragraphs[0] if i == 0 else student_name_tf.add_paragraph()
        p_name.alignment = PP_ALIGN.LEFT
        r_name = p_name.add_run()
        r_name.text = s.name.upper()
        r_name.font.name = FONT
        r_name.font.size = Pt(16)
        r_name.font.bold = False
        r_name.font.color.rgb = BLACK
        set_line_spacing(p_name, 115)
        set_space_after(p_name, 6)

        # Register Number
        p_reg = student_reg_tf.paragraphs[0] if i == 0 else student_reg_tf.add_paragraph()
        p_reg.alignment = PP_ALIGN.LEFT
        r_reg = p_reg.add_run()
        r_reg.text = f"({s.reg_no.upper()})"
        r_reg.font.name = FONT
        r_reg.font.size = Pt(16)
        r_reg.font.bold = False
        r_reg.font.color.rgb = BLACK
        set_line_spacing(p_reg, 115)
        set_space_after(p_reg, 6)

    # ── Right Column: "Under the Guidance of" + Guide Details ──────────────
    guide_txb = slide.shapes.add_textbox(
        Inches(6.2), Inches(3.7), Inches(3.4), Inches(3.0)
    )
    guide_tf = guide_txb.text_frame
    guide_tf.word_wrap = True
    guide_tf.margin_left = Inches(0)
    guide_tf.margin_right = Inches(0)
    guide_tf.margin_top = Inches(0)
    guide_tf.margin_bottom = Inches(0)

    # Line 1: Under the Guidance of
    p_ug = guide_tf.paragraphs[0]
    p_ug.alignment = PP_ALIGN.CENTER
    r_ug = p_ug.add_run()
    r_ug.text = "Under the Guidance of"
    r_ug.font.name = FONT
    r_ug.font.size = Pt(18)
    r_ug.font.bold = False
    r_ug.font.italic = True
    r_ug.font.color.rgb = BLACK
    set_line_spacing(p_ug, 115)
    # Set spacing after to leave exactly one line space gap before guide name
    set_space_after(p_ug, 14)

    # Line 2: Guide Name
    p_gname = guide_tf.add_paragraph()
    p_gname.alignment = PP_ALIGN.CENTER
    r_gname = p_gname.add_run()
    r_gname.text = f"{data.title_slide.guide_name},"
    r_gname.font.name = FONT
    r_gname.font.size = Pt(16)
    r_gname.font.bold = True
    r_gname.font.color.rgb = BLACK
    set_line_spacing(p_gname, 115)

    # Line 3: Designation
    p_gdes = guide_tf.add_paragraph()
    p_gdes.alignment = PP_ALIGN.CENTER
    r_gdes = p_gdes.add_run()
    r_gdes.text = data.title_slide.guide_designation.upper()
    r_gdes.font.name = FONT
    r_gdes.font.size = Pt(14)
    r_gdes.font.bold = True
    r_gdes.font.color.rgb = BLACK
    set_line_spacing(p_gdes, 115)

    # Line 4: Department
    p_gdept = guide_tf.add_paragraph()
    p_gdept.alignment = PP_ALIGN.CENTER
    r_gdept = p_gdept.add_run()
    r_gdept.text = f"Department of {data.title_slide.department}"
    r_gdept.font.name = FONT
    r_gdept.font.size = Pt(12)
    r_gdept.font.bold = False
    r_gdept.font.color.rgb = BLACK
    set_line_spacing(p_gdept, 115)

    logger.info("Built title slide — two-column layout per template spec")
    return slide


# ═══════════════════════════════════════════════════════════════════════════
# ABSTRACT SLIDE
# ═══════════════════════════════════════════════════════════════════════════

def build_abstract_slide(prs, data, slide_num: int):
    """Build the abstract slide."""
    slide = _new_blank_slide(prs)
    add_navy_bar(slide)
    add_footer(slide, data.title_slide.department, slide_num)
    add_heading(slide, "Abstract")
    add_content_box(slide, data.abstract)
    logger.info("Built abstract slide")
    return slide


# ═══════════════════════════════════════════════════════════════════════════
# PROBLEM STATEMENT SLIDE
# ═══════════════════════════════════════════════════════════════════════════

def build_problem_statement_slide(prs, data, slide_num: int):
    """Build the problem statement slide with bullets or paragraph."""
    slide = _new_blank_slide(prs)
    add_navy_bar(slide)
    add_footer(slide, data.title_slide.department, slide_num)
    add_heading(slide, "Problem Statement")

    items = parse_numbered_or_bullet(data.problem_statement)
    if items:
        add_content_with_bullets(slide, items)
    else:
        add_content_box(slide, data.problem_statement)

    logger.info("Built problem statement slide")
    return slide


# ═══════════════════════════════════════════════════════════════════════════
# LITERATURE SURVEY SLIDE (one per entry)
# ═══════════════════════════════════════════════════════════════════════════

def build_literature_slide(prs, data, slide_num: int, entry_index: int):
    """Build one literature survey slide for a single entry."""
    entry = data.literature_survey[entry_index]
    slide = _new_blank_slide(prs)
    add_navy_bar(slide)
    add_footer(slide, data.title_slide.department, slide_num)
    add_heading(slide, "Literature survey")

    txb = slide.shapes.add_textbox(
        CONTENT_LEFT, Inches(1.1),
        CONTENT_WIDTH, Inches(5.5)
    )
    tf = txb.text_frame
    tf.word_wrap = True

    def add_field(tf, label, value, first=False):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        p.alignment = PP_ALIGN.JUSTIFY
        r1 = p.add_run()
        r1.text = label
        r1.font.name = FONT
        r1.font.size = Pt(16)
        r1.font.bold = True
        r1.font.color.rgb = BLACK
        r2 = p.add_run()
        r2.text = value
        r2.font.name = FONT
        r2.font.size = Pt(16)
        r2.font.bold = False
        r2.font.color.rgb = BLACK
        set_line_spacing(p, 115)
        # Blank line after
        pb = tf.add_paragraph()
        set_space_after(pb, 8)

    add_field(tf, f'"{entry.title}"', "", first=True)
    add_field(tf, "Authors: ", entry.author)
    add_field(tf, "Year: ", entry.year)
    add_field(tf, "Topics Covered: ", entry.summary)

    logger.info(f"Built literature slide #{entry_index + 1}")
    return slide


# ═══════════════════════════════════════════════════════════════════════════
# EXISTING SYSTEM SLIDE
# ═══════════════════════════════════════════════════════════════════════════

def build_existing_system_slide(prs, data, slide_num: int):
    """Build the existing system slide."""
    slide = _new_blank_slide(prs)
    add_navy_bar(slide)
    add_footer(slide, data.title_slide.department, slide_num)
    add_heading(slide, "Existing System")

    items = parse_numbered_or_bullet(data.existing_system)
    if items:
        add_content_with_bullets(slide, items)
    else:
        add_content_box(slide, data.existing_system)

    logger.info("Built existing system slide")
    return slide


# ═══════════════════════════════════════════════════════════════════════════
# PROPOSED SYSTEM SLIDE
# ═══════════════════════════════════════════════════════════════════════════

def build_proposed_system_slide(prs, data, slide_num: int):
    """Build the proposed system slide."""
    slide = _new_blank_slide(prs)
    add_navy_bar(slide)
    add_footer(slide, data.title_slide.department, slide_num)
    add_heading(slide, "Proposed System")

    items = parse_numbered_or_bullet(data.proposed_system)
    if items:
        add_content_with_bullets(slide, items)
    else:
        add_content_box(slide, data.proposed_system)

    logger.info("Built proposed system slide")
    return slide


# ═══════════════════════════════════════════════════════════════════════════
# ARCHITECTURE / FLOWCHART SLIDE
# ═══════════════════════════════════════════════════════════════════════════

def build_architecture_slide(prs, data, slide_num: int, image_path: Path):
    """Build the architecture diagram slide."""
    slide = _new_blank_slide(prs)
    add_navy_bar(slide)
    add_footer(slide, data.title_slide.department, slide_num)
    add_heading(slide, "System Architecture")
    add_image_centred(
        slide, image_path, "System Architecture", 1,
        heading_bottom_inches=1.0
    )
    logger.info("Built architecture slide")
    return slide


# ═══════════════════════════════════════════════════════════════════════════
# PROGRESS / OUTPUT IMAGE SLIDE (one per image)
# ═══════════════════════════════════════════════════════════════════════════

def build_image_slide(prs, data, slide_num: int,
                      image_path: Path, caption: str,
                      figure_num: int, heading_text: str):
    """Build a slide with a single centred image and caption."""
    slide = _new_blank_slide(prs)
    add_navy_bar(slide)
    add_footer(slide, data.title_slide.department, slide_num)
    add_heading(slide, heading_text)
    add_image_centred(
        slide, image_path, caption, figure_num,
        heading_bottom_inches=1.0
    )
    logger.info(f"Built image slide: {caption}")
    return slide


# ═══════════════════════════════════════════════════════════════════════════
# DYNAMIC SLIDE BUILDER (Custom / Edited sections)
# ═══════════════════════════════════════════════════════════════════════════

def build_dynamic_slide(prs, data, section, slide_num: int, slide_images: list):
    """
    Build a dynamic slide from custom or edited sections.
    Supports text content (paragraphs or lists) and up to 3 section images.
    """
    logger.info("build_dynamic_slide: section=%s images=%s",
                section.key, [str(p.get("path") if isinstance(p, dict) else p) for p in slide_images])
    for img_info in slide_images:
        img_path = img_info.get("path") if isinstance(img_info, dict) else img_info
        logger.info("  checking image path: %s exists=%s",
                    img_path, Path(img_path).exists() if img_path else False)

    slide = _new_blank_slide(prs)
    add_navy_bar(slide)
    add_footer(slide, data.title_slide.department, slide_num)
    add_heading(slide, section.title)

    has_text = bool(section.content.strip()) if section.content else False
    num_images = len(slide_images)

    if not has_text and num_images == 0:
        logger.info(f"Built empty dynamic slide: {section.title}")
        return slide

    if num_images == 0:
        # Case 1: ONLY TEXT
        if section.format_as_bullets:
            points = content_to_bullet_points(section.content)
            if points:
                items = [("• ", p) for p in points]
                add_content_with_bullets(slide, items)
                logger.info("Slide '%s': rendered as %d bullet points", section.title, len(points))
            else:
                add_content_box(slide, section.content)
                logger.warning("Slide '%s': bullets requested but content cleaned to empty — fell back to paragraph", section.title)
        else:
            items = parse_numbered_or_bullet(section.content)
            if items:
                add_content_with_bullets(slide, items)
                logger.info("Slide '%s': rendered as bullet points (auto-parsed)", section.title)
            else:
                add_content_box(slide, section.content)
                logger.info("Slide '%s': rendered as paragraph", section.title)
        logger.info(f"Built dynamic slide with text only: {section.title}")
        return slide

    if not has_text:
        # Case 2: ONLY IMAGES
        if num_images == 1:
            add_image_centred(slide, slide_images[0]["path"], slide_images[0]["caption"], 1, heading_bottom_inches=1.1)
        elif num_images == 2:
            _add_images_side_by_side(slide, slide_images, top=Inches(1.2), height=Inches(5.0))
        else:
            _add_images_three_grid(slide, slide_images, top=Inches(1.2), height=Inches(5.0))
        logger.info(f"Built dynamic slide with {num_images} images only: {section.title}")
        return slide

    # Case 3: BOTH TEXT AND IMAGES (Split layout: Text on Left, Images on Right)
    left_col_w = Inches(4.3)
    right_col_w = Inches(3.5)
    gap = Inches(0.2)
    left_x = CONTENT_LEFT
    right_x = left_x + left_col_w + gap
    top_y = Inches(1.1)
    height_y = Inches(5.5)

    # Add text in the left column
    if section.format_as_bullets:
        points = content_to_bullet_points(section.content)
        if points:
            items = [("• ", p) for p in points]
            _add_content_with_bullets_custom(slide, items, left=left_x, top=top_y, width=left_col_w, height=height_y)
            logger.info("Slide '%s': split layout rendered as %d bullet points", section.title, len(points))
        else:
            _add_content_box_custom(slide, section.content, left=left_x, top=top_y, width=left_col_w, height=height_y)
            logger.warning("Slide '%s': split layout bullets requested but content cleaned to empty — fell back to paragraph", section.title)
    else:
        items = parse_numbered_or_bullet(section.content)
        if items:
            _add_content_with_bullets_custom(slide, items, left=left_x, top=top_y, width=left_col_w, height=height_y)
            logger.info("Slide '%s': split layout rendered as bullet points (auto-parsed)", section.title)
        else:
            _add_content_box_custom(slide, section.content, left=left_x, top=top_y, width=left_col_w, height=height_y)
            logger.info("Slide '%s': split layout rendered as paragraph", section.title)

    # Add image(s) in the right column
    if num_images == 1:
        _place_single_image_in_col(slide, slide_images[0], right_x, top_y, right_col_w, height_y, 1)
    else:
        _place_stacked_images_in_col(slide, slide_images, right_x, top_y, right_col_w, height_y)

    logger.info(f"Built dynamic slide with text and {num_images} images: {section.title}")
    return slide


def _calculate_fit_dimensions(image_path: Path, max_w, max_h):
    from PIL import Image as PILImage
    try:
        with PILImage.open(str(image_path)) as img:
            aspect = img.height / img.width
    except Exception:
        aspect = 0.75
    
    w = max_w
    h = int(w * aspect)
    if h > max_h:
        h = max_h
        w = int(h / aspect)
    return int(w), int(h)


def _add_images_side_by_side(slide, slide_images, top, height):
    w_max = Inches(3.8)
    gap = Inches(0.4)
    left_start = CONTENT_LEFT
    for idx, img_info in enumerate(slide_images[:2]):
        path = Path(img_info["path"])
        if not path.exists():
            logger.warning("Image path does not exist: %s", path)
            continue
        x = left_start + idx * (w_max + gap)
        w, h = _calculate_fit_dimensions(path, w_max, height - Inches(0.5))
        y = top + (height - Inches(0.5) - h) / 2
        try:
            slide.shapes.add_picture(str(path.resolve()), int(x + (w_max - w) / 2), int(y), w, h)
        except Exception as e:
            logger.error("Failed to add_picture %s: %s", path, e)
            continue
        
        cap_text = f"Figure {idx + 1}: {img_info['caption']}"
        txb = slide.shapes.add_textbox(x, top + height - Inches(0.4), w_max, Inches(0.4))
        tf = txb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = cap_text
        r.font.name = FONT
        r.font.size = Pt(12)


def _add_images_three_grid(slide, slide_images, top, height):
    w_max = Inches(2.4)
    gap = Inches(0.4)
    left_start = CONTENT_LEFT
    for idx, img_info in enumerate(slide_images[:3]):
        path = Path(img_info["path"])
        if not path.exists():
            logger.warning("Image path does not exist: %s", path)
            continue
        x = left_start + idx * (w_max + gap)
        w, h = _calculate_fit_dimensions(path, w_max, height - Inches(0.5))
        y = top + (height - Inches(0.5) - h) / 2
        try:
            slide.shapes.add_picture(str(path.resolve()), int(x + (w_max - w) / 2), int(y), w, h)
        except Exception as e:
            logger.error("Failed to add_picture %s: %s", path, e)
            continue
        
        cap_text = f"Figure {idx + 1}: {img_info['caption']}"
        txb = slide.shapes.add_textbox(x, top + height - Inches(0.4), w_max, Inches(0.4))
        tf = txb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = cap_text
        r.font.name = FONT
        r.font.size = Pt(11)


def _place_single_image_in_col(slide, img_info, left, top, max_w, max_h, fig_num):
    path = Path(img_info["path"])
    if not path.exists():
        logger.warning("Image path does not exist: %s", path)
        return
    w, h = _calculate_fit_dimensions(path, max_w, max_h - Inches(0.5))
    x = left + (max_w - w) / 2
    y = top + (max_h - Inches(0.5) - h) / 2
    try:
        slide.shapes.add_picture(str(path.resolve()), int(x), int(y), w, h)
    except Exception as e:
        logger.error("Failed to add_picture %s: %s", path, e)
        return
    
    cap_text = f"Figure {fig_num}: {img_info['caption']}"
    txb = slide.shapes.add_textbox(left, top + max_h - Inches(0.45), max_w, Inches(0.4))
    tf = txb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = cap_text
    r.font.name = FONT
    r.font.size = Pt(11)


def _place_stacked_images_in_col(slide, slide_images, left, top, max_w, max_h):
    num_imgs = len(slide_images)
    gap = Inches(0.2)
    slot_h = (max_h - (num_imgs - 1) * gap) / num_imgs
    
    for idx, img_info in enumerate(slide_images[:3]):
        path = Path(img_info["path"])
        if not path.exists():
            logger.warning("Image path does not exist: %s", path)
            continue
        slot_top = top + idx * (slot_h + gap)
        img_max_h = slot_h - Inches(0.4)
        if img_max_h < Inches(0.5):
            img_max_h = Inches(0.5)
            
        w, h = _calculate_fit_dimensions(path, max_w, img_max_h)
        x = left + (max_w - w) / 2
        y = slot_top + (img_max_h - h) / 2
        try:
            slide.shapes.add_picture(str(path.resolve()), int(x), int(y), w, h)
        except Exception as e:
            logger.error("Failed to add_picture %s: %s", path, e)
            continue
        
        cap_text = f"Figure {idx + 1}: {img_info['caption']}"
        txb = slide.shapes.add_textbox(left, slot_top + img_max_h + Inches(0.05), max_w, Inches(0.35))
        tf = txb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = cap_text
        r.font.name = FONT
        r.font.size = Pt(9.5)


def _add_content_box_custom(slide, text: str, left, top, width, height, font_size=None, alignment=PP_ALIGN.JUSTIFY):
    if font_size is None:
        font_size = get_body_font_size(text)
        if font_size > 16:
            font_size = 16
    txb = slide.shapes.add_textbox(left, top, width, height)
    tf = txb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = alignment
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(font_size)
    r.font.bold = False
    r.font.color.rgb = BLACK
    set_line_spacing(p, 115)
    return txb


def _add_content_with_bullets_custom(slide, items: list, left, top, width, height):
    txb = slide.shapes.add_textbox(left, top, width, height)
    tf = txb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        bold_part, normal_part = item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.JUSTIFY
        if bold_part:
            r1 = p.add_run()
            r1.text = bold_part
            r1.font.name = FONT
            r1.font.size = Pt(14)
            r1.font.bold = True
            r1.font.color.rgb = BLACK
        if normal_part:
            r2 = p.add_run()
            r2.text = (" " if bold_part else "") + normal_part
            r2.font.name = FONT
            r2.font.size = Pt(14)
            r2.font.bold = False
            r2.font.color.rgb = BLACK
        set_line_spacing(p, 115)
        set_space_after(p, 4)
