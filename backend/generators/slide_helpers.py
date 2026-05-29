"""
Slide Helpers — ALL reusable helper functions for building slides.
Implements the Dr. NGP Institute of Technology college PPT template.
Every slide builder imports from here.
"""

from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
from lxml import etree
from pathlib import Path
from utils.logger import logger

# ── Color constants ───────────────────────────────────────────────────────
NAVY  = RGBColor(0x1F, 0x38, 0x64)
BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# ── Font ──────────────────────────────────────────────────────────────────
FONT = "Times New Roman"

# ── Layout constants (inches) ─────────────────────────────────────────────
CONTENT_LEFT   = Inches(1.38)
CONTENT_TOP    = Inches(0.3)
CONTENT_WIDTH  = Inches(8.02)
SLIDE_W        = Inches(10)
SLIDE_H        = Inches(7.5)
FOOTER_TOP     = Inches(7.1)
FOOTER_H       = Inches(0.35)
BAR_WIDTH      = Inches(1.18)


def get_body_font_size(text: str) -> int:
    """Return 16 for dense text (>80 words), 18 for lighter text."""
    return 16 if len(text.split()) > 80 else 18


def add_navy_bar(slide):
    """Add the left navy bar with rotated college name to any slide."""
    # Navy rectangle
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        left=Inches(0), top=Inches(0),
        width=BAR_WIDTH, height=SLIDE_H
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()

    # Rotated text box — placed so its centre aligns with the bar centre.
    # When a shape rotates 270° in OOXML the pivot is the shape centre, so
    # we swap width/height and position the textbox so its centre equals the
    # bar centre (bar_cx = BAR_WIDTH/2, bar_cy = SLIDE_H/2).
    # textbox width (after swap) = SLIDE_H, height = BAR_WIDTH
    # left = bar_cx - SLIDE_H/2,  top = bar_cy - BAR_WIDTH/2
    txb_left = int(BAR_WIDTH / 2 - SLIDE_H / 2)
    txb_top  = int(SLIDE_H  / 2 - BAR_WIDTH / 2)
    txb = slide.shapes.add_textbox(
        txb_left, txb_top,
        SLIDE_H, BAR_WIDTH          # width=SLIDE_H, height=BAR_WIDTH (swapped)
    )
    txb.rotation = 270.0
    tf = txb.text_frame
    tf.word_wrap = False
    tf.auto_size = None
    # Vertical anchor: middle — centres text top-to-bottom inside the box
    txb_sp = txb._element
    txBody = txb_sp.find(qn('p:txBody'))
    if txBody is None:
        txBody = txb_sp.find('.//{http://schemas.openxmlformats.org/presentationml/2006/main}txBody')
    if txBody is not None:
        bodyPr = txBody.find(qn('a:bodyPr'))
        if bodyPr is not None:
            bodyPr.set('anchor', 'ctr')
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Dr. NGP INSTITUTE OF TECHNOLOGY"
    r.font.name = FONT
    r.font.size = Pt(27.5)
    r.font.bold = False
    r.font.color.rgb = WHITE


def add_footer(slide, department: str, page_num: int):
    """Add footer bar with department name (centre) and page number (right)."""
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(1.18), FOOTER_TOP,
        SLIDE_W, FOOTER_TOP,
    )
    line.line.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)
    line.line.width = Pt(0.5)

    # Department name — centred
    txb_c = slide.shapes.add_textbox(
        CONTENT_LEFT, FOOTER_TOP, CONTENT_WIDTH, FOOTER_H
    )
    tf = txb_c.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = f"Department of {department}"
    r.font.name = FONT
    r.font.size = Pt(15.5)
    r.font.bold = True
    r.font.color.rgb = BLACK

    # Page number — right-aligned
    txb_r = slide.shapes.add_textbox(
        CONTENT_LEFT, FOOTER_TOP, CONTENT_WIDTH, FOOTER_H
    )
    tf2 = txb_r.text_frame
    tf2.word_wrap = False
    p2 = tf2.paragraphs[0]
    p2.alignment = PP_ALIGN.RIGHT
    r2 = p2.add_run()
    r2.text = str(page_num)
    r2.font.name = FONT
    r2.font.size = Pt(14)
    r2.font.bold = True
    r2.font.color.rgb = BLACK


def add_heading(slide, text: str):
    """Add the slide main heading (24pt Bold Left)."""
    txb = slide.shapes.add_textbox(
        CONTENT_LEFT, Inches(0.3), CONTENT_WIDTH, Inches(0.55)
    )
    tf = txb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(24)
    r.font.bold = True
    r.font.color.rgb = BLACK


def add_heading_centre(slide, text, left, top, width, size=24):
    """Add a centred heading at custom position."""
    txb = slide.shapes.add_textbox(left, top, width, Inches(0.55))
    tf = txb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = BLACK


def add_content_box(slide, text: str, top=Inches(1.0), height=Inches(5.9),
                    font_size: int = None, alignment=PP_ALIGN.JUSTIFY):
    """Add a content text box with auto font size."""
    if font_size is None:
        font_size = get_body_font_size(text)
    txb = slide.shapes.add_textbox(CONTENT_LEFT, top, CONTENT_WIDTH, height)
    tf = txb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = alignment
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(font_size)
    r.font.bold = False
    r.font.color.rgb = BLACK
    set_line_spacing(p, 115)
    return txb


def add_content_with_bullets(slide, items: list, top=Inches(1.0)):
    """
    Add content as bullet points or numbered items.
    items = list of (bold_part, normal_part) tuples.
    If bold_part is empty, render as plain bullet.
    """
    txb = slide.shapes.add_textbox(CONTENT_LEFT, top, CONTENT_WIDTH, Inches(5.9))
    tf = txb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        bold_part, normal_part = item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.JUSTIFY
        if bold_part:
            r1 = p.add_run()
            r1.text = bold_part
            r1.font.name = FONT
            r1.font.size = Pt(16)
            r1.font.bold = True
            r1.font.color.rgb = BLACK
        if normal_part:
            r2 = p.add_run()
            r2.text = (" " if bold_part else "") + normal_part
            r2.font.name = FONT
            r2.font.size = Pt(16)
            r2.font.bold = False
            r2.font.color.rgb = BLACK
        set_line_spacing(p, 115)
        set_space_after(p, 6)


def add_styled_textbox(slide, left, top, width, height, paragraphs_data):
    """Add a text box with multiple styled paragraphs."""
    txb = slide.shapes.add_textbox(left, top, width, height)
    tf = txb.text_frame
    tf.word_wrap = True
    for i, pd in enumerate(paragraphs_data):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = pd.get("alignment", PP_ALIGN.LEFT)
        r = p.add_run()
        r.text = pd["text"]
        r.font.name = FONT
        r.font.size = Pt(pd.get("size", 16))
        r.font.bold = pd.get("bold", False)
        r.font.italic = pd.get("italic", False)
        r.font.color.rgb = pd.get("color", BLACK)
        if pd.get("spacing"):
            set_line_spacing(p, pd["spacing"])


def add_image_centred(slide, image_path: Path, caption: str,
                      figure_num: int, heading_bottom_inches=1.0):
    """Add a centred image with figure caption below."""
    from PIL import Image as PILImage

    p = Path(image_path)
    if not p.exists():
        logger.warning("Image path does not exist, skipping: %s", p)
        return

    MAX_W = Inches(7.5)
    MAX_H = Inches(5.4)
    top = Inches(heading_bottom_inches)

    try:
        with PILImage.open(str(p.resolve())) as img:
            aspect = img.height / img.width
    except Exception as e:
        logger.error("Failed to open image %s: %s", p, e)
        return

    img_w = MAX_W
    img_h = int(img_w * aspect)
    if img_h > MAX_H:
        img_h = MAX_H
        img_w = int(img_h / aspect)

    img_left = int((SLIDE_W - img_w) / 2)
    try:
        slide.shapes.add_picture(str(p.resolve()), img_left, top, img_w, img_h)
    except Exception as e:
        logger.error("Failed to add_picture %s: %s", p, e)
        return

    # Caption below image
    cap_top = top + img_h + Inches(0.1)
    cap_text = f"Figure {figure_num}: {caption}"
    txb = slide.shapes.add_textbox(CONTENT_LEFT, cap_top, CONTENT_WIDTH, Inches(0.35))
    tf = txb.text_frame
    p_cap = tf.paragraphs[0]
    p_cap.alignment = PP_ALIGN.CENTER
    r = p_cap.add_run()
    r.text = cap_text
    r.font.name = FONT
    r.font.size = Pt(14)
    r.font.bold = False
    r.font.color.rgb = BLACK


def set_line_spacing(paragraph, percent: int):
    """Set line spacing as a percentage (e.g. 115 = 1.15x)."""
    pPr = paragraph._p.get_or_add_pPr()
    lnSpc = pPr.find(qn("a:lnSpc"))
    if lnSpc is None:
        lnSpc = etree.SubElement(pPr, qn("a:lnSpc"))
    for child in list(lnSpc):
        lnSpc.remove(child)
    spcPct = etree.SubElement(lnSpc, qn("a:spcPct"))
    spcPct.set("val", str(percent * 1000))


def set_space_after(paragraph, pts: int):
    """Set space after paragraph in points."""
    pPr = paragraph._p.get_or_add_pPr()
    spcAft = pPr.find(qn("a:spcAft"))
    if spcAft is None:
        spcAft = etree.SubElement(pPr, qn("a:spcAft"))
    for child in list(spcAft):
        spcAft.remove(child)
    spcPts = etree.SubElement(spcAft, qn("a:spcPts"))
    spcPts.set("val", str(pts * 100))


def set_slide_background_white(slide):
    """Set slide background to pure white."""
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE


def parse_numbered_or_bullet(text: str):
    """
    Parse text into list of (bold_part, normal_part) tuples.
    Detects patterns like:
      "1. Title\\nDescription"
      "• Item text"
      "- Item text"
    Returns empty list if plain paragraph.
    """
    import re
    lines = text.strip().split('\n')
    items = []
    numbered_pattern = re.compile(r'^(\d+\.\s*.+?)$')
    bullet_pattern = re.compile(r'^[•\-\*]\s*(.+)$')

    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        # Numbered with sub-description
        nm = numbered_pattern.match(line)
        if nm:
            bold = line
            desc = ""
            if (i + 1 < len(lines) and lines[i + 1].strip() and
                    not numbered_pattern.match(lines[i + 1].strip())):
                desc = lines[i + 1].strip()
                i += 1
            items.append((bold, desc))
        else:
            bm = bullet_pattern.match(line)
            if bm:
                items.append(("• ", bm.group(1)))
            else:
                # Not a list — return empty to use plain text
                return []
        i += 1
    return items if len(items) > 1 else []
